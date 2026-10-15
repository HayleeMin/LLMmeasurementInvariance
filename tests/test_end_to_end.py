"""End-to-end test of the pilot pipeline on synthetic data with a mock embedder (no network, no GPU).

Run from the project root:  python tests/test_end_to_end.py
It copies the project to a temp folder, generates fake SemEval and pilot data, and runs every step.
"""
import gzip
import random
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np

SRC = Path(__file__).resolve().parents[1]
tmp = Path(tempfile.mkdtemp()) / "pilot"
shutil.copytree(SRC, tmp, ignore=shutil.ignore_patterns("data", "outputs", "__pycache__"))
(tmp / "data").mkdir()
(tmp / "outputs" / "cache").mkdir(parents=True)
sys.path.insert(0, str(tmp))

from openpyxl import load_workbook  # noqa: E402

from src.common import load_config, write_jsonl  # noqa: E402
from src.embedder import MockEmbedder  # noqa: E402

cfg = load_config()
R = random.Random(1)

# ---------- fake SemEval: 10 targets; context cue "alpha" in period 1, "beta" share in period 2 = grade
sem = tmp / cfg["gate1_semeval"]["data_dir"]
targets = [f"word{i}_nn" for i in range(10)]
grades = {t: i / 9 for i, t in enumerate(targets)}
(sem / "truth").mkdir(parents=True)
(sem / "targets.txt").write_text("\n".join(targets))
(sem / "truth" / "graded.txt").write_text("\n".join(f"{t}\t{g:.3f}" for t, g in grades.items()))
for period in (1, 2):
    for kind in ("lemma", "token"):
        (sem / f"corpus{period}" / kind).mkdir(parents=True)
    L, T = [], []
    for t in targets:
        for k in range(60):
            cue = "beta" if (period == 2 and R.random() < grades[t]) else "alpha"
            w = t.split("_")[0]
            toks = ["the", w, "was", cue, f"n{k}", "today"]
            T.append(" ".join(toks))
            L.append(" ".join(["the", t, "be", cue, f"n{k}", "today"]))
    for kind, lines in (("lemma", L), ("token", T)):
        with gzip.open(sem / f"corpus{period}" / kind / f"c{period}.txt.gz", "wt") as f:
            f.write("\n".join(lines))

mock_sem = lambda spec: MockEmbedder(n_layers=4, dim=32, cues={"a": ["alpha"], "b": ["beta"]}, seed=0)  # noqa: E731
from scripts.s01_gate1_semeval import run as run_g1  # noqa: E402

g1 = run_g1(mock_sem, cfg["models"][:2], cfg, sem, tmp / "outputs" / "gate1_results.csv")
assert g1.spearman_drift.max() > 0.8, g1
print("Gate 1 OK: max rho", round(g1.spearman_drift.max(), 3))

# ---------- fake pilot corpus with category cues
cues = {"exhaustion": ["drained", "exhausted"], "cynicism": ["stopped caring", "cynical"],
        "reduced_efficacy": ["incompetent", "mistakes"], "depressive_symptoms": ["hopeless", "no joy"],
        "ordinary_tiredness": ["a bit", "nap"], "other_none": ["rate", "program"]}
docs, truth = [], {}
for i in range(1500):
    cat = R.choice(list(cues))
    y = R.randint(2012, 2022)
    txt = f"I think {R.choice(['burnout', 'burned out', 'burnt out'])} {R.choice(cues[cat])} here. Next sentence {i}."
    docs.append({"id": str(i), "author": f"u{R.randint(1, 300)}", "thread": R.randint(1, 200), "year": y, "text": txt})
    truth[str(i)] = cat
write_jsonl(tmp / "data" / "pilot" / "burnout_comments.jsonl", docs)
write_jsonl(tmp / "data" / "pilot" / "background.jsonl", [{"text": f"random comment {i}"} for i in range(500)])
(tmp / "data" / "pilot" / "year_totals.csv").write_text("year,comments\n" + "\n".join(f"{y},1000000" for y in range(2012, 2023)))

subprocess.run([sys.executable, "scripts/s03_make_coding_sheets.py"], cwd=tmp, check=True)

# fake coders: truth with 10% / 15% noise
import pandas as pd  # noqa: E402

from src.common import CATEGORY_LABELS  # noqa: E402

key = pd.read_csv(tmp / "data" / "pilot" / "coding_key.csv", dtype={"doc_id": str})
for coder, noise in (("A", 0.10), ("B", 0.15)):
    wb = load_workbook(tmp / "materials" / f"Coding_Sheet_Coder{coder}.xlsx")
    ws = wb["Code"]
    rr = random.Random(coder)
    for r in range(4, ws.max_row + 1):
        iid = ws.cell(r, 1).value
        doc = key.set_index("item_id").loc[iid, "doc_id"]
        cat = truth[doc] if rr.random() > noise else rr.choice(list(cues))
        ws.cell(r, 3).value = CATEGORY_LABELS[cat]
    wb.save(tmp / f"coder{coder}.xlsx")
subprocess.run([sys.executable, "scripts/s04_agreement.py", "check", "coderA.xlsx", "coderB.xlsx"], cwd=tmp, check=True)
wb = load_workbook(tmp / "materials" / "Disagreements.xlsx")
ws = wb["Resolve"]
for r in range(2, ws.max_row + 1):
    iid = ws.cell(r, 1).value
    if iid:
        ws.cell(r, 5).value = CATEGORY_LABELS[truth[key.set_index("item_id").loc[iid, "doc_id"]]]
wb.save(tmp / "materials" / "Disagreements.xlsx")
subprocess.run([sys.executable, "scripts/s04_agreement.py", "finalize", "coderA.xlsx", "coderB.xlsx",
                "materials/Disagreements.xlsx"], cwd=tmp, check=True)

# ---------- Gate 2 with a mock embedder that knows the category cues (anchors contain cue words too)
anchor_cues = {k: v for k, v in cues.items() if k != "other_none"}
anchor_cues["exhaustion"] += ["drained", "exhausted", "energy", "tank", "depleted", "worn"]
anchor_cues["cynicism"] += ["caring", "cynical", "detached", "bitter", "indifference", "checked out"]
anchor_cues["reduced_efficacy"] += ["competent", "mistakes", "accomplish", "useless", "unproductive", "failing"]
anchor_cues["depressive_symptoms"] += ["hopeless", "joy", "empty", "sad", "worthless", "despair"]
anchor_cues["ordinary_tiredness"] += ["a bit", "nap", "mini", "little", "slight", "weekend"]
mock_pilot = lambda spec: MockEmbedder(n_layers=4, dim=32, cues=anchor_cues, seed=3)  # noqa: E731
from scripts.s05_gate2_classify import run as run_g2  # noqa: E402

labels = pd.read_csv(tmp / "data" / "pilot" / "final_labels.csv")
g2, confs = run_g2(mock_pilot, cfg["models"][:2], cfg, labels, ["bg text"] * 50)
g2.to_csv(tmp / "outputs" / "gate2_results.csv", index=False)
import json  # noqa: E402

json.dump(confs, open(tmp / "outputs" / "gate2_confusion.json", "w"))
print("Gate 2 best macro-F1:", round(g2.macro_f1.max(), 3), "ECE:", round(g2.loc[g2.macro_f1.idxmax(), "ece"], 3))
assert g2.macro_f1.max() > 0.5, g2

subprocess.run([sys.executable, "scripts/s06_select_model.py"], cwd=tmp, check=True)

from scripts.s08_power_inputs import all_usage_inputs, coded_sample_inputs  # noqa: E402
from src.common import read_jsonl  # noqa: E402

shares, iccs = coded_sample_inputs(labels)
sel = json.load(open(tmp / "outputs" / "selected_model.json")) or {"model": "mock", "layer": 4}
extra = all_usage_inputs(mock_pilot(None), sel, labels, read_jsonl(tmp / "data" / "pilot" / "usages.jsonl"),
                         ["bg"] * 20, cfg)
print("Power inputs OK: null drift SD", round(extra["null_drift_sd"], 4), "| coded shares", {k: round(v, 2) for k, v in shares.items()})

# ---------- control screen logic (no network)
from scripts.s07_control_screen import finalize, screen_row  # noqa: E402

docs_by_year = {2019: ["We discussed process today. Nothing else."] * 50,
                2021: ["Remote process is hard. Zoom all day."] * 25 + ["Our process works. Fine."] * 25}
docs_by_year.update({y: ["process here. ok."] * 40 for y in range(2012, 2023) if y not in docs_by_year})
row = screen_row("process", docs_by_year, {y: 1e6 for y in range(2012, 2023)}, cfg)
assert row["fails_remote_screen"], row
row2 = screen_row("budget", {y: ["budget talk. ok."] * 40 for y in range(2012, 2023)}, {y: 1e6 for y in range(2012, 2023)}, cfg)
assert not row2["fails_remote_screen"]
scr, widened = finalize(pd.DataFrame([row, row2]), np.log10(40), cfg)
assert list(scr.in_final_set) == [False, True], scr
print("Control screen OK")
print("\nALL TESTS PASSED in", tmp)
