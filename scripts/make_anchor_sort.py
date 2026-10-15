"""Build the blind expert anchor-sort workbook and its answer key.

Usage: python scripts/make_anchor_sort.py
Outputs: materials/Expert_Anchor_Sort.xlsx (send to experts), materials/anchor_key.csv (keep private)
"""
import csv
import random
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
cfg = yaml.safe_load(open(ROOT / "config.yaml"))
anchors = yaml.safe_load(open(ROOT / "materials" / "anchors.yaml"))

rows = [(k, v["label"], s) for k, v in anchors.items() for s in v["anchors"]]
random.Random(cfg["seed"]).shuffle(rows)
labels = [v["label"] for v in anchors.values()]

wb = Workbook()
ins = wb.active
ins.title = "Instructions"
text = [
    "Anchor sentence sort: burnout components",
    "",
    "Each sentence below uses the word 'burnout' (or 'burned out' / 'burnt out').",
    "For each sentence, choose from the drop-down list the ONE component the word most clearly conveys.",
    "Work alone, and don't look anything up. It takes about 15 minutes.",
    "",
    "Components:",
]
for v in anchors.values():
    text.append(f"  - {v['label']}: {v['definition']}")
text += ["", "Enter your name in cell B1 of the 'Sort' sheet, then return the file."]
for i, t in enumerate(text, 1):
    ins.cell(row=i, column=1, value=t)
ins["A1"].font = Font(bold=True, size=13)
ins.column_dimensions["A"].width = 140

ws = wb.create_sheet("Sort")
ws["A1"], ws["A1"].font = "Expert name:", Font(bold=True)
ws["B1"].fill = PatternFill("solid", fgColor="FFF2CC")
hdr = ["ID", "Sentence", "Your choice"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=3, column=j, value=h)
    c.font = Font(bold=True)
dv = DataValidation(type="list", formula1='"' + ",".join(labels) + '"', allow_blank=True)
ws.add_data_validation(dv)
key = []
for i, (k, lab, s) in enumerate(rows, 1):
    aid = f"A{i:02d}"
    ws.cell(row=3 + i, column=1, value=aid)
    c = ws.cell(row=3 + i, column=2, value=s)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ch = ws.cell(row=3 + i, column=3)
    ch.fill = PatternFill("solid", fgColor="FFF2CC")
    dv.add(ch)
    key.append((aid, k, lab, s))
ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 95
ws.column_dimensions["C"].width = 32
ws.freeze_panes = "A4"
wb.save(ROOT / "materials" / "Expert_Anchor_Sort.xlsx")

with open(ROOT / "materials" / "anchor_key.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["id", "component", "label", "sentence"])
    w.writerows(key)
print(f"Wrote {len(rows)} anchors to Expert_Anchor_Sort.xlsx and anchor_key.csv")
