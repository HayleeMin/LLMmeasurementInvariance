"""Intercoder agreement and label resolution.

Step 1 (after both coders finish):
    python scripts/s04_agreement.py check  CoderA.xlsx CoderB.xlsx
      -> Krippendorff's alpha (overall; exhaustion-vs-tiredness subset), outputs/agreement.json,
         materials/Disagreements.xlsx (fill the 'resolved' column after discussion)
Step 2 (after resolving disagreements):
    python scripts/s04_agreement.py finalize CoderA.xlsx CoderB.xlsx materials/Disagreements.xlsx
      -> data/pilot/final_labels.csv
"""
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.common import CATEGORY_LABELS, DATA, LABEL_TO_KEY, MATERIALS, OUT, load_config  # noqa: E402
from src.stats import krippendorff_alpha_nominal  # noqa: E402


def read_codes(path):
    ws = load_workbook(path, data_only=True)["Code"]
    rows = []
    for r in range(4, ws.max_row + 1):
        iid = ws.cell(r, 1).value
        if iid:
            rows.append({"item_id": iid, "label": ws.cell(r, 3).value, "excerpt": ws.cell(r, 2).value,
                         "mixed": (ws.cell(r, 4).value or "").strip().upper() == "Y",
                         "unsure": (ws.cell(r, 5).value or "").strip().upper() == "Y"})
    return pd.DataFrame(rows)


def merged(pa, pb):
    a, b = read_codes(pa), read_codes(pb)
    m = a.merge(b[["item_id", "label", "mixed", "unsure"]], on="item_id", suffixes=("_a", "_b"))
    missing = m[m.label_a.isna() | m.label_b.isna()]
    if len(missing):
        sys.exit(f"{len(missing)} items are uncoded by at least one coder, e.g. {missing.item_id.head().tolist()}")
    m["key_a"], m["key_b"] = m.label_a.map(LABEL_TO_KEY), m.label_b.map(LABEL_TO_KEY)
    return m


if __name__ == "__main__":
    mode, pa, pb = sys.argv[1], sys.argv[2], sys.argv[3]
    cfg = load_config()
    m = merged(pa, pb)
    if mode == "check":
        alpha = krippendorff_alpha_nominal(m.key_a, m.key_b)
        sub = m[m.key_a.isin(["exhaustion", "ordinary_tiredness"]) | m.key_b.isin(["exhaustion", "ordinary_tiredness"])]
        alpha_et = krippendorff_alpha_nominal(sub.key_a, sub.key_b)
        res = {"n_items": len(m), "alpha_overall": alpha, "n_exh_tired_items": len(sub),
               "alpha_exhaustion_vs_tiredness": alpha_et, "raw_agreement": float((m.key_a == m.key_b).mean()),
               "threshold": cfg["coding"]["min_alpha"],
               "passes": bool(alpha >= cfg["coding"]["min_alpha"]),
               "merge_exhaustion_tiredness": bool(alpha_et < cfg["coding"]["min_alpha"]),
               "pct_mixed": float((m.mixed_a | m.mixed_b).mean()), "pct_unsure": float((m.unsure_a | m.unsure_b).mean()),
               "crosstab": pd.crosstab(m.key_a, m.key_b).to_dict()}
        OUT.mkdir(exist_ok=True)
        json.dump(res, open(OUT / "agreement.json", "w"), indent=2)
        print(json.dumps({k: v for k, v in res.items() if k != "crosstab"}, indent=2))
        d = m[m.key_a != m.key_b]
        wb = Workbook()
        ws = wb.active
        ws.title = "Resolve"
        for j, h in enumerate(["item_id", "excerpt", "coder_A", "coder_B", "resolved"], 1):
            ws.cell(row=1, column=j, value=h)
        dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORY_LABELS.values()) + '"')
        ws.add_data_validation(dv)
        for i, r in enumerate(d.itertuples(), 2):
            for j, v in enumerate([r.item_id, r.excerpt, r.label_a, r.label_b], 1):
                ws.cell(row=i, column=j, value=v)
            dv.add(ws.cell(row=i, column=5))
        ws.column_dimensions["B"].width = 90
        wb.save(MATERIALS / "Disagreements.xlsx")
        print(f"{len(d)} disagreements written to materials/Disagreements.xlsx")
    elif mode == "finalize":
        res = {}
        if len(sys.argv) > 4:
            ws = load_workbook(sys.argv[4], data_only=True)["Resolve"]
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, 1).value:
                    res[ws.cell(r, 1).value] = ws.cell(r, 5).value
        m["final"] = [ka if ka == kb else LABEL_TO_KEY.get(res.get(i)) for i, ka, kb in zip(m.item_id, m.key_a, m.key_b)]
        if m.final.isna().any():
            sys.exit(f"{m.final.isna().sum()} disagreements still unresolved in Disagreements.xlsx")
        key = pd.read_csv(DATA / "pilot" / "coding_key.csv")
        out = key.merge(m[["item_id", "final", "mixed_a", "mixed_b"]], on="item_id")
        out.to_csv(DATA / "pilot" / "final_labels.csv", index=False)
        print(out.final.value_counts().to_string())
