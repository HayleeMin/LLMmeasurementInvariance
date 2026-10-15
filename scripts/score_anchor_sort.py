"""Score returned expert anchor sorts against the key (80% rule).

Usage: python scripts/score_anchor_sort.py returned/*.xlsx
Writes outputs/anchor_sort_results.csv and prints anchors to revise.
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
key = pd.read_csv(ROOT / "materials" / "anchor_key.csv")
files = [Path(p) for p in sys.argv[1:]]
if not files:
    sys.exit("Pass the returned expert workbooks as arguments.")

answers = {}
for f in files:
    ws = load_workbook(f, data_only=True)["Sort"]
    name = ws["B1"].value or f.stem
    for r in range(4, ws.max_row + 1):
        aid, choice = ws.cell(r, 1).value, ws.cell(r, 3).value
        if aid:
            answers.setdefault(aid, {})[name] = choice

out = key.copy()
n_experts = len(files)
out["n_answered"] = [sum(v is not None for v in answers.get(a, {}).values()) for a in out["id"]]
out["n_correct"] = [sum(v == lab for v in answers.get(a, {}).values()) for a, lab in zip(out["id"], out["label"])]
out["pct_correct"] = out["n_correct"] / out["n_answered"].clip(lower=1)
out["revise"] = out["pct_correct"] < 0.80
(ROOT / "outputs").mkdir(exist_ok=True)
out.to_csv(ROOT / "outputs" / "anchor_sort_results.csv", index=False)

print(f"Experts: {n_experts}. Overall placement agreement: {out['n_correct'].sum() / out['n_answered'].sum():.1%}")
print(out.groupby("label")["pct_correct"].mean().round(3).to_string())
bad = out[out["revise"]]
print(f"\nAnchors to revise (<80% placed correctly): {len(bad)}")
for _, r in bad.iterrows():
    print(f"  {r['id']} [{r['label']}, {r['pct_correct']:.0%}] {r['sentence']}")
