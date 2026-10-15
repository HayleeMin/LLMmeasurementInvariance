"""Extract burnout usages from the pilot corpus, draw the year-stratified coding sample, and write
blind coding sheets for two coders.

Usage:  python scripts/s03_make_coding_sheets.py
Output: data/pilot/usages.jsonl (all usages), data/pilot/coding_key.csv (private: item -> year, ids)
        materials/Coding_Sheet_CoderA.xlsx, materials/Coding_Sheet_CoderB.xlsx (no dates shown)
"""
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.common import (CATEGORY_LABELS, DATA, MATERIALS, burnout_pattern, clean_html,  # noqa: E402
                        extract_usages, load_config, read_jsonl, write_jsonl)

GUIDE = [
    "Code what the word 'burnout' (marked >>like this<<) conveys about a person's state in this excerpt. One category per item.",
    "1. No person's state described (rates, programs, the word as a topic, generic advice, engine/rocket/tyre senses) -> Other / none.",
    "2. Exhaustion vs. ordinary tiredness: chronic (weeks+ or recurring), work-attributed, or impairing recovery -> Exhaustion;",
    "   short-lived, situational, expected to pass with rest, or hedged as mild ('a bit', 'mini') -> Ordinary tiredness.",
    "   Duration unclear: strong depletion language ('drained', 'nothing left') -> Exhaustion; otherwise Ordinary tiredness.",
    "3. Low mood, hopelessness, or loss of joy extending beyond work, or depression named -> Depressive symptoms;",
    "   negativity aimed at the job, organization, or people served -> Cynicism / depersonalization.",
    "4. Competence, mistakes, productivity, or lost sense of accomplishment -> Reduced professional efficacy.",
    "5. Two components equally prominent: code the one stated first and put Y in 'mixed'.",
    "6. Unsure: give your best code and put Y in 'unsure'. Do not skip items. Do not discuss items until both coders finish.",
]


def load_docs(cfg):
    pc = cfg["pilot_corpus"]
    if pc["source"] == "csv":
        df = pd.read_csv(pc["csv_path"])
        df["year"] = pd.to_datetime(df["date"]).dt.year
        return [{"id": str(r.id), "author": r.author, "thread": r.thread, "year": int(r.year),
                 "text": clean_html(str(r.text))} for r in df.itertuples()]
    return read_jsonl(DATA / "pilot" / "burnout_comments.jsonl")


def write_sheet(path, items):
    wb = Workbook()
    ins = wb.active
    ins.title = "Guide"
    for i, t in enumerate(GUIDE, 1):
        ins.cell(row=i, column=1, value=t)
    ins.column_dimensions["A"].width = 150
    ws = wb.create_sheet("Code")
    ws["A1"], ws["B1"] = "Coder name:", ""
    ws["A1"].font = Font(bold=True)
    hdr = ["item_id", "excerpt", "category", "mixed (Y)", "unsure (Y)", "notes"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=3, column=j, value=h).font = Font(bold=True)
    dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORY_LABELS.values()) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, it in enumerate(items, 4):
        ws.cell(row=i, column=1, value=it["item_id"])
        c = ws.cell(row=i, column=2, value=it["display"])
        c.alignment = Alignment(wrap_text=True, vertical="top")
        cat = ws.cell(row=i, column=3)
        cat.fill = PatternFill("solid", fgColor="FFF2CC")
        dv.add(cat)
    for col, w in zip("ABCDEF", [9, 95, 30, 10, 10, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    wb.save(path)


if __name__ == "__main__":
    cfg = load_config()
    pc = cfg["pilot_corpus"]
    pat = burnout_pattern(cfg)
    usages = []
    for d in load_docs(cfg):
        for k, (text, s, e) in enumerate(extract_usages(d["text"], pat)):
            usages.append({"usage_id": f"{d['id']}_{k}", "doc_id": d["id"], "author": d.get("author"),
                           "thread": d.get("thread"), "year": d["year"], "text": text, "start": s, "end": e})
    write_jsonl(DATA / "pilot" / "usages.jsonl", usages)
    df = pd.DataFrame(usages)
    df = df[(df.year >= pc["years"][0]) & (df.year <= pc["years"][1])]
    print("Usages per year:\n" + df.groupby("year").size().to_string())

    # Year-stratified sample: at most one usage per comment; min per year, remainder proportional.
    r = np.random.default_rng(cfg["seed"])
    one = df.sample(frac=1, random_state=cfg["seed"]).drop_duplicates("doc_id")
    n_total, n_min = pc["coding_sample_size"], pc["coding_min_per_year"]
    counts = one.groupby("year").size()
    alloc = np.minimum(counts, n_min)
    rest = n_total - alloc.sum()
    extra = np.floor((counts - alloc) / max((counts - alloc).sum(), 1) * rest).astype(int)
    alloc = np.minimum(counts, alloc + extra)
    years = list(counts.index)
    while alloc.sum() < min(n_total, counts.sum()):  # top up rounding shortfall, one per year in random order
        for y in r.permutation(years):
            if alloc.sum() >= n_total:
                break
            if alloc[y] < counts[y]:
                alloc[y] += 1
    picked = pd.concat([g.sample(int(alloc[y]), random_state=cfg["seed"]) for y, g in one.groupby("year")])
    picked = picked.sample(frac=1, random_state=cfg["seed"] + 1).reset_index(drop=True)
    picked["item_id"] = [f"P{i:03d}" for i in range(1, len(picked) + 1)]
    picked["display"] = [t[:s] + ">>" + t[s:e] + "<<" + t[e:] for t, s, e in zip(picked.text, picked.start, picked.end)]
    picked[["item_id", "usage_id", "doc_id", "author", "thread", "year", "text", "start", "end"]].to_csv(
        DATA / "pilot" / "coding_key.csv", index=False)
    items = picked[["item_id", "display"]].to_dict("records")
    write_sheet(MATERIALS / "Coding_Sheet_CoderA.xlsx", items)
    write_sheet(MATERIALS / "Coding_Sheet_CoderB.xlsx", items)
    print(f"Coding sample: {len(picked)} usages; per year:\n" + picked.groupby("year").size().to_string())
